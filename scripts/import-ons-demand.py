"""Extract published ONS observations without estimating suppressed values.

Usage: python import-ons-demand.py WORKBOOK OUTPUT_JSON
Only reads the workbook. Keep its SHA-256 with the release evidence.
"""
import sys, json, hashlib, datetime
from pathlib import Path
import openpyxl
from openpyxl.utils import get_column_letter

source = Path(sys.argv[1])
destination = Path(sys.argv[2])
wb = openpyxl.load_workbook(source, read_only=True, data_only=True)
areas = {'Lincoln':{'code':'E07000138','name':'Lincoln'}, 'Newark and Sherwood':{'code':'E07000175','name':'Newark and Sherwood'}}
records=[]
def observation(value, period, sheet, row, col):
    numeric = isinstance(value,(int,float)) and not isinstance(value,bool) and value >= 0 and int(value)==value
    marker = str(value).strip() if value is not None else None
    status = 'available' if numeric else 'suppressed' if marker and 'x' in marker.lower() else 'unavailable'
    return {'period':period,'value':int(value) if numeric else None,'status':status,'sourceMarker':None if numeric else marker,'sourceCell':f'{sheet}!{get_column_letter(col)}{row}'}

sheet=wb['Table 5']
header=list(next(sheet.iter_rows(min_row=5,max_row=5,values_only=True)))
assert header[:6]==['Region','ITL2','Local Authority District','SOC 4 digit code','SOC 4 digit label','Local Authority Code']
assert header[-1]=='2026Q2'
periods=list(enumerate(header[-4:],start=len(header)-3))
for rownum,row in enumerate(sheet.iter_rows(min_row=6,values_only=True),start=6):
    if row[0]!='East Midlands': break
    if row[2] not in areas: continue
    area=areas[row[2]]
    assert row[5]==area['code']
    soc=str(row[3])
    if soc == 'Unknown': continue
    assert len(soc)==4 and soc.isdigit(), f'Unexpected classification: {rownum} {row[:6]}'
    values=[observation(row[col-1],period,sheet.title,rownum,col) for col,period in periods]
    records.append({'geography':area['code'],'soc2020':soc,'occupationTitle':row[4],'frequency':'quarterly','observations':values})

sheet=wb['Table 3'];header=list(next(sheet.iter_rows(min_row=5,max_row=5,values_only=True)))
assert header[:3]==['Region','SOC 4 digit code','SOC 4 digit label']
assert header[-1]=='Jul-26'
periods=list(enumerate(header[-6:],start=len(header)-5))
found=False
for rownum,row in enumerate(sheet.iter_rows(min_row=6,values_only=True),start=6):
    if row[0]!='East Midlands':
        if found: break
        continue
    found=True;soc=str(row[1])
    if soc == 'Unknown': continue
    assert len(soc)==4 and soc.isdigit(), f'Unexpected classification: {rownum} {row[:3]}'
    values=[observation(row[col-1],datetime.datetime.strptime(period,'%b-%y').strftime('%Y-%m'),sheet.title,rownum,col) for col,period in periods]
    records.append({'geography':'E12000004','soc2020':soc,'occupationTitle':row[2],'frequency':'monthly','observations':values})

assert all(sum(row['geography']==code for row in records)>=400 for code in ['E07000138','E07000175','E12000004'])
assert len({(row['geography'],row['soc2020']) for row in records})==len(records), 'Ambiguous geography/SOC key'
dataset={'schemaVersion':1,'provider':'Office for National Statistics','upstreamProvider':'Textkernel',
  'dataset':'Labour demand volumes by Standard Occupation Classification (SOC 2020), UK',
  'sourceUrl':'https://www.ons.gov.uk/employmentandlabourmarket/peopleinwork/employmentandemployeetypes/datasets/labourdemandvolumesbystandardoccupationclassificationsoc2020uk',
  'downloadUrl':'https://www.ons.gov.uk/file?uri=/employmentandlabourmarket/peopleinwork/employmentandemployeetypes/datasets/labourdemandvolumesbystandardoccupationclassificationsoc2020uk/january2017tojuly2026/labourdemandbyoccupation.xlsx',
  'edition':'January 2017 to July 2026','publishedAt':'2026-08-21','retrievedAt':datetime.datetime.now(datetime.timezone.utc).isoformat(),
  'sourceFile':{'name':source.name,'sha256':hashlib.sha256(source.read_bytes()).hexdigest()},
  'classification':'SOC2020','measure':'new-online-job-adverts','unit':'adverts',
  'geographies':[*areas.values(),{'code':'E12000004','name':'East Midlands'}],
  'licence':{'name':'Open Government Licence v3.0','url':'https://www.nationalarchives.gov.uk/doc/open-government-licence/version/3/'},
  'attribution':'Source: Office for National Statistics, using Textkernel online job adverts data. Crown copyright.',
  'qualityNote':'Official statistics in development. ONS reports increased uncertainty since November 2025 and advises caution with short-term changes. Some 2025 and 2026 observations are suppressed after changes in source coverage.',
  'coverageNote':'Published new online adverts for the wider SOC 2020 occupation group, not live vacancies, unique jobs, a forecast or a measure of all hiring. Local areas use LAD 2023 boundaries; East Midlands is a region. Quarterly district observations and monthly regional observations are different periods and must not be compared as like-for-like counts.',
  'records':records}
destination.write_text(json.dumps(dataset,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
for area in dataset['geographies']:
    selected=[row for row in records if row['geography']==area['code']]
    counts={status:sum(row['observations'][-1]['status']==status for row in selected) for status in ['available','suppressed','unavailable']}
    all_codes={row['soc2020'] for row in records}
    print(area['name'],len(selected),counts,'Missing groups:', sorted(all_codes-{row['soc2020'] for row in selected}))
print(json.dumps([r for r in records if r['soc2020']=='2134'],ensure_ascii=False))
print('Cover notes:')
for row in wb['Cover sheet'].iter_rows(min_row=19,values_only=True):
    if row[0]: print(row[0])
wb.close()
