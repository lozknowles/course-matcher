"""Extract published 2025 ASHE annual medians and their quality markers.

Usage: python import-ashe-pay.py SOURCES_DIRECTORY RETRIEVED_AT
SOURCES_DIRECTORY must contain the official Table 14 and Table 15 ZIP downloads.
Requires openpyxl for reading only. Original workbooks are never modified.
"""
from collections import Counter
from datetime import datetime
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile
import hashlib
import json
import math
import re
import sys

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
BASE_URL = 'https://www.ons.gov.uk/employmentandlabourmarket/peopleinwork/earningsandworkinghours/datasets/'


def number(value):
    return isinstance(value, (int, float)) and not isinstance(value, bool) and math.isfinite(value)


def extract(directory, table, retrieved_at):
    regional = table == 15
    archive = directory / f'ashetable{table}2025provisional.zip'
    raw = archive.read_bytes()
    books, source_files = {}, []
    with ZipFile(BytesIO(raw)) as zipped:
        for kind in ('a', 'b'):
            matches = [name for name in zipped.namelist() if name.endswith('.xlsx')
                       and '(4)' in name and f'.7{kind} ' in name and 'Annual pay - Gross' in name]
            if len(matches) != 1:
                raise ValueError(f'Expected exactly one annual gross workbook: table {table}, {kind}')
            name = matches[0]
            data = zipped.read(name)
            source_files.append({'name': name, 'sha256': hashlib.sha256(data).hexdigest()})
            books[kind] = openpyxl.load_workbook(BytesIO(data), read_only=True, data_only=True)

    sheets = {kind: book['Full-Time'] for kind, book in books.items()}
    for sheet in sheets.values():
        assert '2025' in sheet['A1'].value and 'full-time employee jobs' in sheet['A1'].value
        assert sheet['B5'].value == 'Code' and sheet['D5'].value == 'Median'
    quality_rows = {row[0].row: row for row in sheets['b'].iter_rows(min_row=6, max_col=4)}
    records = []
    for row in sheets['a'].iter_rows(min_row=6, max_col=4):
        title, code, _, median = [cell.value for cell in row]
        if not re.fullmatch(r'\d{4}', str(code)):
            continue
        if regional and not str(title).strip().startswith('East Midlands,'):
            continue
        quality_row = quality_rows[row[0].row]
        assert quality_row[0].value == title and quality_row[1].value == code
        cv = quality_row[3].value
        if median in ('x', '..') or (number(cv) and cv > 20) or cv == 'x':
            status, value, flag = 'suppressed', None, 'suppressed'
            note = 'ONS has withheld this median for reliability or disclosure reasons.'
        elif number(median) and median >= 0 and number(cv) and 0 <= cv <= 20:
            status, value = 'available', median
            flag = 'precise' if cv <= 5 else 'reasonably-precise' if cv <= 10 else 'use-with-caution'
            note = f'ONS coefficient of variation: {cv:g}%.'
            if cv > 10:
                note += ' Use with caution.'
        elif number(median) and median >= 0 and cv == '.':
            # The dot refers to the CV, not the separately published median.
            status, value, flag = 'available', median, 'cv-unavailable'
            note = 'ONS publishes this median, but its numeric quality score (coefficient of variation) is unavailable. Use with caution.'
        else:
            status, value, flag = 'unavailable', None, 'not-published'
            note = 'ONS has not published a usable median for this occupation group.'
        records.append({
            'soc2020': str(code),
            'occupationTitle': str(title).strip().removeprefix('East Midlands,').strip() if regional else str(title).strip(),
            'value': value, 'status': status, 'qualityFlag': flag, 'qualityNote': note,
            'coefficientOfVariation': cv if number(cv) else None,
            'sourceCell': f'Full-Time!D{row[0].row}', 'qualitySourceCell': f'Full-Time!D{row[0].row}',
            'sourceMarker': median if isinstance(median, str) else None,
            'qualityMarker': cv if isinstance(cv, str) else None,
        })
    assert len(records) == len({r['soc2020'] for r in records}) == 412
    geography = {'code': 'E12000004', 'name': 'East Midlands'} if regional else {'code': 'K02000001', 'name': 'United Kingdom'}
    result = {
        'schemaVersion': 1, 'provider': 'Office for National Statistics',
        'dataset': 'ASHE Table 15 (4).7a / 7b' if regional else 'ASHE Table 14.7a / 7b',
        'sourceUrl': BASE_URL + ('regionbyoccupation4digitsoc2010ashetable15' if regional else 'occupation4digitsoc2010ashetable14'),
        'retrievedAt': retrieved_at, 'year': 2025, 'edition': '2025 provisional', 'geography': geography,
        'measure': {'id': 'median-gross-annual-full-time', 'label': 'Median gross annual pay · full-time employees', 'unit': 'GBP/year'},
        'classification': 'SOC2020',
        'licence': {'name': 'Open Government Licence v3.0', 'url': 'https://www.nationalarchives.gov.uk/doc/open-government-licence/version/3/'},
        'coverageNote': f"Workplace-based {geography['name']} occupation groups. Full-time employee jobs on adult rates, in the same job for more than a year; tax year ended 5 April 2025. Excludes self-employment. Not starting salaries or current advertised pay.",
        'qualityNote': 'A dot in the CV workbook means that the quality score is unavailable. A separately published numeric median is retained with a caution. Suppressed or blank medians are never estimated.',
        'sourceFiles': source_files, 'archiveSha256': hashlib.sha256(raw).hexdigest(), 'records': records,
    }
    for book in books.values():
        book.close()
    return result


def main():
    directory, retrieved_at = Path(sys.argv[1]), sys.argv[2]
    assert datetime.fromisoformat(retrieved_at).tzinfo is not None
    for table, filename in ((15, 'ons-pay-reference.json'), (14, 'ons-uk-pay-reference.json')):
        data = extract(directory, table, retrieved_at)
        (ROOT / 'student-hub' / filename).write_text(json.dumps(data, indent=2, ensure_ascii=False) + '\n', encoding='utf-8')
        print(filename, dict(Counter(r['status'] for r in data['records'])))


if __name__ == '__main__':
    main()
